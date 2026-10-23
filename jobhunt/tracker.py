"""Export the pipeline to output/tracker.xlsx and dashboard_jobs.json (one row per job)."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .store import Store

COLUMNS = ["found_on", "match_score", "ats_score", "title", "company", "location", "source",
           "easy_apply", "apply_url", "matched_keywords", "missing_keywords", "resume_path",
           "contacts", "status"]


def export(store: Store, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in store.jobs():
        contacts = "; ".join(
            f"{c['name']} ({c['title']}) {c['email'] or c['linkedin_url']}" for c in store.contacts(row["key"]))
        ws.append([row["found_on"], row["match_score"], row["ats_score"], row["title"], row["company"],
                   row["location"], row["source"], "Yes" if row["easy_apply"] else "No",
                   row["apply_url"] or row["url"], row["matched_keywords"], row["missing_keywords"],
                   row["resume_path"], contacts, row["status"]])
    widths = [11, 8, 8, 40, 28, 22, 10, 9, 50, 40, 40, 40, 60, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_json(store: Store, path: Path) -> Path:
    """Rows for the dashboard. No recruiter names or emails: only counts."""
    rows = []
    for row in store.jobs():
        contacts = store.contacts(row["key"])
        rows.append({
            "key": row["key"], "found_on": row["found_on"], "title": row["title"],
            "company": row["company"], "location": row["location"], "source": row["source"],
            "apply_type": "Easy Apply" if row["easy_apply"] else "Company site",
            "apply_url": row["apply_url"] or row["url"], "fit": row["match_score"],
            "ats_before": row["ats_before"], "ats_after": row["ats_score"],
            "contacts": len(contacts), "drafts": sum(1 for c in contacts if c["draft_id"]),
            "status": row["status"],
        })
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(rows, indent=1, ensure_ascii=False))
    return path
