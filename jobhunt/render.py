"""Render an ATS-friendly DOCX: one column, standard section headings, real text
(no tables, text boxes, images, icons, or ligature-mangled fonts)."""
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def build_master(master: dict):
    """The master resume as-is, in the same shape as a tailor.Tailored result."""
    skills = [s for group in (master.get("skills") or {}).values() for s in group]
    return SimpleNamespace(
        summary=master.get("summary", ""),
        experience=[SimpleNamespace(id=e["id"], bullets=e["bullets"]) for e in master["experience"]],
        skills=skills,
        projects=[SimpleNamespace(**p) for p in master.get("projects", [])],
    )


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)


def _line(doc, left: str, right: str = "", bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(left)
    r.bold = bold
    if right:
        p.add_run(f"  |  {right}")


def _bullet(doc, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)


def build(master: dict, tailored) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(50)
        section.top_margin = section.bottom_margin = Pt(40)

    c = master["contact"]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(master["name"])
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph(" | ".join(x for x in (c.get("email"), c.get("phone"), c.get("location"),
                                                 c.get("linkedin")) if x))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _heading(doc, "Professional Summary")
    doc.add_paragraph(tailored.summary)

    _heading(doc, "Professional Experience")
    bullets = {r.id: r.bullets for r in tailored.experience}
    for e in master["experience"]:
        if e["id"] not in bullets:
            continue
        _line(doc, f"{e['title']}, {e['company']}", f"{e['start']} – {e['end']}", bold=True)
        for b in bullets[e["id"]]:
            _bullet(doc, b)

    _heading(doc, "Skills")
    doc.add_paragraph(", ".join(tailored.skills))

    if tailored.projects:
        _heading(doc, "Projects")
        for pr in tailored.projects:
            _line(doc, pr.name, bold=True)
            doc.add_paragraph(pr.description)

    _heading(doc, "Education")
    for ed in master.get("education", []):
        _line(doc, f"{ed['degree']}, {ed['school']}", ed.get("dates", ""))

    if master.get("certifications"):
        _heading(doc, "Certifications")
        for cert in master["certifications"]:
            _bullet(doc, cert)

    if master.get("leadership"):
        _heading(doc, "Leadership")
        for item in master["leadership"]:
            _bullet(doc, item)

    if master.get("languages"):
        _heading(doc, "Languages")
        doc.add_paragraph(", ".join(master["languages"]))
    return doc


def plain_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def save_docx(doc: Document, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
