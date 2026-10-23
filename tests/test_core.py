import copy
from pathlib import Path

import pytest
import yaml
from docx import Document

from jobhunt import match, render, resolve, tailor
from jobhunt.config import load_config
from jobhunt.models import Job
from jobhunt.sources import adzuna, jsearch
from jobhunt.store import Store

ROOT = Path(__file__).resolve().parent.parent
MASTER = yaml.safe_load((ROOT / "examples" / "master_resume.example.yaml").read_text())
MASTER["experience"][0]["bullets"] = [
    "Managed SG, MY and IN payroll for 5 global clients with 100% statutory compliance.",
    "Resolved Payroll Control Center (PCC) incidents and performed root cause analysis.",
    "Used LSMW and schema debugging (PE01/PE02) for data integrity.",
]
MASTER["skills"]["technical"] = ["SAP HCM Payroll", "HRIS", "Workday", "Microsoft Excel"]

JD = """We are hiring an SAP HCM Payroll Consultant in Hyderabad. 2-4 years of experience in
SAP HR / payroll, Payroll Control Center, PCRs and schemas, LSMW, SuccessFactors Employee
Central is a plus. Strong stakeholder management. Experience with regression testing."""


@pytest.fixture
def cfg(tmp_path):
    c = load_config()
    c["verified_extra_skills"] = []
    c["paths"]["database"] = tmp_path / "t.db"
    c["paths"]["output_dir"] = tmp_path / "out"
    return c


# --- match ---
@pytest.mark.parametrize("text,ok", [
    ("Requires 2-4 years of SAP HCM", True),
    ("3 to 6 yrs experience", True),
    ("Minimum 8 years of payroll experience", False),
    ("10-15 years in HR leadership", False),
    ("5+ years in Workday", False),
    ("1-3 years", True),
    ("Fresher role, no experience needed", True),
])
def test_experience_filter(text, ok):
    assert match.experience_ok(text, (2, 4)) is ok


@pytest.mark.parametrize("loc,ok", [
    ("Hyderabad, Telangana, India", True), ("Remote", True), ("Bengaluru, India", True),
    ("London, UK", False), ("Singapore", False),
])
def test_location_filter(loc, ok):
    assert match.location_ok(loc) is ok


def test_keywords_word_boundaries():
    assert "LSMW" in match.keywords_in("used lsmw daily")
    assert "SQL" not in match.keywords_in("nosqlish")  # no partial-word hits
    assert "Payroll Control Center" in match.keywords_in("PCC alerts")


def test_score_prefers_relevant_jobs():
    resume = tailor.master_text(MASTER)
    good = match.score(JD, resume, "SAP HCM Payroll Consultant", ["SAP HCM"])
    bad = match.score("Java Spring Boot microservices developer, Kubernetes", resume, "Java Developer", ["SAP HCM"])
    assert good["score"] > 50 > bad["score"]
    assert "SuccessFactors" in good["missing"]


# --- sources ---
def test_jsearch_parse_prefers_direct_company_link():
    payload = {"data": [{
        "job_title": "HRIS Analyst", "employer_name": "Acme", "employer_website": "https://www.acme.com",
        "job_city": "Hyderabad", "job_country": "IN", "job_is_remote": False,
        "job_publisher": "LinkedIn", "job_apply_link": "https://www.linkedin.com/jobs/view/1",
        "apply_options": [{"publisher": "LinkedIn", "apply_link": "https://www.linkedin.com/jobs/view/1", "is_direct": False},
                          {"publisher": "Acme Careers", "apply_link": "https://careers.acme.com/j/9", "is_direct": True}],
        "job_description": "HRIS role",
    }]}
    [job] = jsearch.parse(payload)
    assert job.apply_url == "https://careers.acme.com/j/9"
    assert not job.easy_apply
    assert job.company_domain == "acme.com"


def test_adzuna_parse():
    [job] = adzuna.parse({"results": [{"title": "Payroll Analyst", "company": {"display_name": "Beta"},
                                       "location": {"display_name": "Hyderabad"},
                                       "redirect_url": "https://adzuna.in/x", "description": "payroll"}]})
    assert (job.title, job.company, job.source) == ("Payroll Analyst", "Beta", "adzuna")


# --- resolve ---
def test_employer_domain():
    assert resolve.employer_domain("https://careers.deloitte.com/job/1") == "deloitte.com"
    assert resolve.employer_domain("https://boards.greenhouse.io/acme/jobs/1") == ""
    assert resolve.employer_domain("https://acme.wd1.myworkdayjobs.com/x") == ""


def test_extract_description_from_jsonld():
    html = """<html><script type="application/ld+json">{"@type":"JobPosting",
      "description":"<p>Configure <b>SAP Payroll</b> schemas</p>"}</script><body>nav junk</body></html>"""
    assert "Configure" in resolve.extract_description(html)
    assert "nav junk" not in resolve.extract_description(html)


# --- store ---
def test_store_dedupes_across_sources(tmp_path):
    s = Store(tmp_path / "db.sqlite")
    a = Job("HRIS Analyst", "Acme", "Hyderabad", "u1", "jsearch")
    b = Job("HRIS  analyst", "ACME", "Hyderabad", "u2", "linkedin")
    assert s.add_job(a) and not s.add_job(b)
    assert s.used("jobs") == 1


# --- tailoring guard rails ---
def test_filter_skills_drops_invented_skills():
    allowed = tailor.allowed_skills(MASTER, {"verified_extra_skills": []})
    out = tailor.filter_skills(["SAP HR", "SAP HCM Payroll", "SuccessFactors", "Workday HCM Configuration",
                                "HRIS"], allowed, JD)
    assert "SAP HCM Payroll" in out and "HRIS" in out
    assert "SAP HR" in out                      # JD's spelling of a skill she has
    assert "SuccessFactors" not in out          # not on her resume
    assert "Workday HCM Configuration" not in out  # overstates "Workday"


def test_unsupported_numbers_flags_invented_metrics():
    assert tailor.unsupported_numbers("Cut payroll errors by 40% for 5 clients at 100%", MASTER) == ["40%"]


def test_ats_match():
    r = tailor.ats_match(["LSMW", "PCRs", "Payroll Control Center"], "Used LSMW and Payroll Control Center")
    assert r["percent"] == 67 and r["missing"] == ["PCRs"]


def test_tailor_end_to_end_with_mocked_claude(cfg, monkeypatch, tmp_path):
    fake = tailor.Tailored(
        jd_keywords=["SAP HR", "Payroll Control Center", "LSMW", "SuccessFactors", "regression testing"],
        summary="SAP HR payroll consultant with 2 years of APAC payroll support.",
        experience=[tailor.RoleBullets(id="acme", bullets=[
            "Resolved Payroll Control Center (PCC) incidents with root cause analysis.",
            "Used LSMW and payroll schema debugging (PE01/PE02); ran regression testing.",
            "Improved accuracy by 37%.",  # invented metric -> must be flagged
        ]), tailor.RoleBullets(id="made-up-job", bullets=["Led SAP team"])],
        skills=["SAP HR", "SAP HCM Payroll", "SuccessFactors", "HRIS"],
        projects=[tailor.Project(name="Payroll migration", description="Migrated payroll.")],
        gaps=["SuccessFactors Employee Central"],
    )
    monkeypatch.setattr(tailor.llm, "structured", lambda *a, **k: copy.deepcopy(fake))
    job = Job("SAP HCM Payroll Consultant", "Acme", "Hyderabad", "u", "jsearch", description=JD)
    rep = tailor.tailor(job, MASTER, cfg, tmp_path)

    text = "\n".join(p.text for p in Document(rep["resume_path"]).paragraphs)
    assert "Led SAP team" not in text            # unknown role dropped
    assert "SuccessFactors" not in text.split("SKILLS")[1].split("PROJECTS")[0]
    assert "Acme Corp" in text and "Jan 2024" in text  # employer/dates copied from master
    assert rep["review_numbers"] == ["37%"]
    assert rep["ats_after"] > rep["ats_before"]
    assert "SuccessFactors" in rep["missing"]    # honest: not claimed, so not matched


def test_render_is_ats_safe(tmp_path):
    doc = render.build(MASTER, render.build_master(MASTER))
    assert not doc.tables and not doc.inline_shapes
    text = render.plain_text(doc)
    for heading in ("PROFESSIONAL SUMMARY", "PROFESSIONAL EXPERIENCE", "SKILLS", "EDUCATION"):
        assert heading in text
